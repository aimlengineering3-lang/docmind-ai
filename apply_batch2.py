"""Batch 2: M1 fixes (header/footer stripping, DOCX locators, validation) + M2 (embeddings, FAISS, FTS5, hybrid retrieval).
Run from INSIDE the docmind-ai folder:  python apply_batch2.py
"""
from pathlib import Path

if not Path("docmind").is_dir():
    raise SystemExit("Run this from inside the docmind-ai folder (the one that contains docmind/).")

FILES = {}
FILES['docmind/parsing.py'] = r'''import hashlib
import re
import unicodedata
from collections import Counter
from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .schemas import Segment

# A PDF page with fewer extractable characters than this is treated as scanned.
MIN_CHARS_FOR_TEXT_PAGE = 30
# DOCX text without any headings is cut into segments of about this many words,
# so citations can still point to a paragraph range instead of just the file.
MAX_HEADINGLESS_WORDS = 500


class ParseError(Exception):
    """Raised when a file cannot be parsed (corrupt, encrypted, unsupported)."""


def file_id(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()[:12]


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _norm_for_repeat(paragraph: str) -> str:
    # "Page 3 of 10" and "Page 4 of 10" must count as the same footer.
    return re.sub(r"\d+", "#", paragraph.lower()).strip()


def _strip_repeated(
    pages: list[list[str]],
    min_pages: int = 3,
    ratio: float = 0.5,
    max_len: int = 120,
    edge: int = 2,
) -> list[list[str]]:
    """Drop running headers/footers: short paragraphs near the top or bottom of a
    page whose normalised text shows up on at least `ratio` of all pages."""
    if len(pages) < min_pages:
        return pages

    def edge_idx(n: int) -> set[int]:
        return set(range(min(edge, n))) | set(range(max(n - edge, 0), n))

    counts: Counter[str] = Counter()
    for paras in pages:
        counts.update(
            {_norm_for_repeat(paras[i]) for i in edge_idx(len(paras)) if len(paras[i]) <= max_len}
        )
    repeated = {key for key, c in counts.items() if c >= ratio * len(pages)}

    cleaned = []
    for paras in pages:
        idx = edge_idx(len(paras))
        cleaned.append(
            [
                p
                for i, p in enumerate(paras)
                if not (i in idx and len(p) <= max_len and _norm_for_repeat(p) in repeated)
            ]
        )
    return cleaned


def parse_pdf(path: Path, doc_id: str) -> list[Segment]:
    try:
        pdf = pymupdf.open(path)
    except Exception as e:
        raise ParseError(f"Cannot open PDF '{path.name}': {e}") from e

    with pdf:
        if pdf.needs_pass:
            raise ParseError(f"PDF '{path.name}' is password-protected")

        pages: list[list[str]] = []
        for page in pdf:
            # Each text block (type 0) is roughly a paragraph; sort=True gives reading order.
            blocks = page.get_text("blocks", sort=True)
            paragraphs = [" ".join(b[4].split()) for b in blocks if b[6] == 0]
            pages.append([p for p in paragraphs if p])

    segments = []
    for page_no, paragraphs in enumerate(_strip_repeated(pages), start=1):
        text = clean_text("\n\n".join(paragraphs))
        segments.append(
            Segment(
                doc_id=doc_id,
                doc_name=path.name,
                page=page_no,
                text=text,
                needs_ocr=len(text) < MIN_CHARS_FOR_TEXT_PAGE,
            )
        )
    return segments


def _iter_blocks(doc):
    """Yield paragraphs and tables in document order (python-docx doesn't do this natively)."""
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _table_to_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            value = " ".join(cell.text.split())
            if not cells or value != cells[-1]:  # merged cells repeat; collapse them
                cells.append(value)
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def parse_docx(path: Path, doc_id: str) -> list[Segment]:
    try:
        doc = Document(str(path))
    except Exception as e:
        raise ParseError(f"Cannot open DOCX '{path.name}': {e}") from e

    segments: list[Segment] = []
    heading: str | None = None
    buffer: list[str] = []
    buf_words = 0
    para_no = 0  # counts non-empty paragraphs and tables
    start_para = 1

    def flush() -> None:
        nonlocal buf_words, start_para
        text = clean_text("\n\n".join(buffer))
        if text:
            # Before the first heading (or in heading-less files) cite a paragraph range.
            section = heading or f"paras {start_para}-{para_no}"
            segments.append(
                Segment(doc_id=doc_id, doc_name=path.name, section=section, text=text)
            )
        buffer.clear()
        buf_words = 0
        start_para = para_no + 1

    for block in _iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = block.style.name if block.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                flush()
                heading = text
            para_no += 1
            buffer.append(text)
            buf_words += len(text.split())
        else:
            table_text = _table_to_text(block)
            if not table_text.strip():
                continue
            para_no += 1
            buffer.append(table_text)
            buf_words += len(table_text.split())
        if heading is None and buf_words >= MAX_HEADINGLESS_WORDS:
            flush()
    flush()
    return segments


PARSERS = {".pdf": parse_pdf, ".docx": parse_docx}


def parse_document(path: str | Path) -> list[Segment]:
    path = Path(path)
    if not path.is_file():
        raise ParseError(f"File not found: {path}")
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ParseError(f"Unsupported file type: '{path.suffix}'")
    return parser(path, file_id(path))
'''
FILES['docmind/chunking.py'] = r'''import re

from .schemas import Chunk, Segment


def _split_long(words: list[str], max_words: int, overlap: int) -> list[str]:
    """Sliding window for a single paragraph that exceeds max_words."""
    step = max_words - overlap
    windows = []
    for start in range(0, len(words), step):
        windows.append(" ".join(words[start : start + max_words]))
        if start + max_words >= len(words):
            break
    return windows


def chunk_segment(
    text: str, max_words: int = 180, overlap_words: int = 30
) -> list[str]:
    if overlap_words >= max_words:
        raise ValueError("overlap_words must be smaller than max_words")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]

    units: list[str] = []
    for p in paragraphs:
        words = p.split()
        if len(words) > max_words:
            units.extend(_split_long(words, max_words, overlap_words))
        else:
            units.append(p)

    chunks: list[str] = []
    buffer: list[str] = []
    buffer_words = 0
    for unit in units:
        n = len(unit.split())
        if buffer and buffer_words + n > max_words:
            chunks.append("\n\n".join(buffer))
            buffer, buffer_words = [], 0
        buffer.append(unit)
        buffer_words += n
    if buffer:
        chunks.append("\n\n".join(buffer))
    return chunks


def chunk_segments(
    segments: list[Segment],
    max_words: int = 180,
    overlap_words: int = 30,
    min_words: int = 5,
) -> list[Chunk]:
    """Chunk each segment independently so a chunk never spans pages/sections."""
    chunks: list[Chunk] = []
    for seg in segments:
        if not seg.text:
            continue  # e.g. scanned page awaiting OCR
        for text in chunk_segment(seg.text, max_words, overlap_words):
            word_count = len(text.split())
            if word_count < min_words:
                continue  # page numbers, stray labels
            idx = len(chunks)
            chunks.append(
                Chunk(
                    chunk_id=f"{seg.doc_id}:{idx:04d}",
                    doc_id=seg.doc_id,
                    doc_name=seg.doc_name,
                    page=seg.page,
                    section=seg.section,
                    chunk_index=idx,
                    text=text,
                    word_count=word_count,
                    method=seg.method,
                )
            )
    return chunks'''
FILES['docmind/embeddings.py'] = r'''import logging

import numpy as np

log = logging.getLogger("docmind.embeddings")


class BgeEmbedder:
    """BAAI/bge-small-en-v1.5 on CPU. Vectors are L2-normalised, so inner product = cosine.

    The model loads lazily, so commands that never embed (list, delete) start instantly.
    """

    # bge-v1.5 retrieval convention: prefix the *query* only, never the passages.
    QUERY_PREFIX = "Represent this sentence for searching relevant passages: "

    def __init__(self, model_name: str = "BAAI/bge-small-en-v1.5", batch_size: int = 16):
        self.model_name = model_name
        self.batch_size = batch_size
        self._model = None

    @property
    def model(self):
        if self._model is None:
            from sentence_transformers import SentenceTransformer

            log.info("Loading embedding model %s (first run downloads ~130 MB)", self.model_name)
            self._model = SentenceTransformer(self.model_name, device="cpu")
        return self._model

    def embed_passages(self, texts: list[str]) -> np.ndarray:
        vecs = self.model.encode(
            texts,
            batch_size=self.batch_size,
            normalize_embeddings=True,
            convert_to_numpy=True,
            show_progress_bar=len(texts) > 64,
        )
        return vecs.astype(np.float32)

    def embed_query(self, query: str) -> np.ndarray:
        vec = self.model.encode(
            [self.QUERY_PREFIX + query], normalize_embeddings=True, convert_to_numpy=True
        )
        return vec.astype(np.float32)
'''
FILES['docmind/index.py'] = r'''import argparse
import logging
import os
import re
import sqlite3
from collections import defaultdict
from pathlib import Path
from typing import Literal, Protocol

import faiss
import numpy as np
from pydantic import BaseModel

from .ingest import ingest_file
from .parsing import ParseError
from .schemas import Chunk

log = logging.getLogger("docmind.index")

DEFAULT_DB = os.environ.get("DOCMIND_DB", "data/docmind.db")
RRF_K = 60  # standard reciprocal-rank-fusion constant

SCHEMA = """
CREATE TABLE IF NOT EXISTS documents (
    doc_id   TEXT PRIMARY KEY,
    doc_name TEXT NOT NULL,
    n_chunks INTEGER NOT NULL,
    added_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS chunks (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    chunk_id    TEXT NOT NULL UNIQUE,
    doc_id      TEXT NOT NULL,
    doc_name    TEXT NOT NULL,
    page        INTEGER,
    section     TEXT,
    chunk_index INTEGER NOT NULL,
    text        TEXT NOT NULL,
    word_count  INTEGER NOT NULL,
    method      TEXT NOT NULL,
    embedding   BLOB NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_chunks_doc ON chunks(doc_id);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
    text, content='chunks', content_rowid='id', tokenize='porter unicode61'
);
CREATE TRIGGER IF NOT EXISTS chunks_ai AFTER INSERT ON chunks BEGIN
    INSERT INTO chunks_fts(rowid, text) VALUES (new.id, new.text);
END;
CREATE TRIGGER IF NOT EXISTS chunks_ad AFTER DELETE ON chunks BEGIN
    INSERT INTO chunks_fts(chunks_fts, rowid, text) VALUES ('delete', old.id, old.text);
END;
"""

CHUNK_COLUMNS = "id, chunk_id, doc_id, doc_name, page, section, chunk_index, text, word_count, method"

# Dropped from the keyword query; they only add noise to an OR-style BM25 match.
STOPWORDS = frozenset(
    "a an and are as at be by did do does for from how in is it of on or that the this to "
    "was were what when where which who with".split()
)

Mode = Literal["hybrid", "semantic", "keyword"]


class Embedder(Protocol):
    def embed_passages(self, texts: list[str]) -> np.ndarray: ...
    def embed_query(self, query: str) -> np.ndarray: ...


class Hit(BaseModel):
    chunk: Chunk
    score: float  # ranking score: RRF (hybrid), cosine (semantic) or -bm25 (keyword)
    semantic_score: float | None = None  # cosine similarity, filled for every hit unless mode=keyword
    keyword_score: float | None = None  # -bm25, only if the chunk matched the keyword query


class Index:
    """SQLite is the single source of truth (chunks, FTS5 keyword index, embeddings).
    FAISS is an in-memory exact index rebuilt from SQLite, so deletes and re-indexing
    never leave the two stores out of sync."""

    def __init__(self, embedder: Embedder, db_path: str | Path = DEFAULT_DB):
        self.embedder = embedder
        if str(db_path) != ":memory:":
            Path(db_path).parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(str(db_path))
        self.db.row_factory = sqlite3.Row
        self.db.executescript(SCHEMA)
        self._faiss = None
        self._row_ids: list[int] = []
        self._row_docs: list[str] = []
        self._rebuild()

    # ---------- write path ----------
    def _rebuild(self) -> None:
        rows = self.db.execute("SELECT id, doc_id, embedding FROM chunks ORDER BY id").fetchall()
        if not rows:
            self._faiss, self._row_ids, self._row_docs = None, [], []
            return
        mat = np.vstack([np.frombuffer(r["embedding"], dtype=np.float32) for r in rows])
        self._faiss = faiss.IndexFlatIP(mat.shape[1])
        self._faiss.add(np.ascontiguousarray(mat))
        self._row_ids = [r["id"] for r in rows]
        self._row_docs = [r["doc_id"] for r in rows]

    def has_document(self, doc_id: str) -> bool:
        return self.db.execute("SELECT 1 FROM documents WHERE doc_id=?", (doc_id,)).fetchone() is not None

    def add_chunks(self, chunks: list[Chunk]) -> None:
        """Index the chunks of ONE document."""
        if not chunks:
            return
        vecs = self.embedder.embed_passages([c.text for c in chunks]).astype(np.float32)
        first = chunks[0]
        with self.db:  # single transaction: documents + chunks + FTS (via trigger)
            self.db.execute(
                "INSERT OR REPLACE INTO documents(doc_id, doc_name, n_chunks) VALUES (?,?,?)",
                (first.doc_id, first.doc_name, len(chunks)),
            )
            self.db.executemany(
                "INSERT INTO chunks(chunk_id, doc_id, doc_name, page, section, chunk_index,"
                " text, word_count, method, embedding) VALUES (?,?,?,?,?,?,?,?,?,?)",
                [
                    (c.chunk_id, c.doc_id, c.doc_name, c.page, c.section, c.chunk_index,
                     c.text, c.word_count, c.method, v.tobytes())
                    for c, v in zip(chunks, vecs)
                ],
            )
        self._rebuild()

    def add_file(self, path: str | Path) -> tuple[str, int]:
        """Parse, chunk, embed and index a file. Returns (doc_id, chunks_added);
        chunks_added is 0 when the document was already indexed or has no extractable text."""
        segments, chunks = ingest_file(path)
        if not segments:
            raise ParseError(f"No content found in '{Path(path).name}'")
        doc_id = segments[0].doc_id
        if self.has_document(doc_id):
            return doc_id, 0
        if not chunks:
            log.warning("'%s' has no extractable text (scanned? OCR comes later)", Path(path).name)
            return doc_id, 0
        self.add_chunks(chunks)
        return doc_id, len(chunks)

    def delete_document(self, doc_id: str) -> bool:
        with self.db:
            cur = self.db.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
            self.db.execute("DELETE FROM documents WHERE doc_id=?", (doc_id,))
        if cur.rowcount:
            self._rebuild()
        return cur.rowcount > 0

    def documents(self) -> list[sqlite3.Row]:
        return self.db.execute(
            "SELECT doc_id, doc_name, n_chunks, added_at FROM documents ORDER BY added_at, doc_name"
        ).fetchall()

    # ---------- read path ----------
    def _semantic(self, qvec: np.ndarray, n: int, doc_ids: set[str] | None) -> list[tuple[int, float]]:
        if self._faiss is None:
            return []
        # Flat index scores every row anyway, so with a doc filter we score all rows and filter after.
        fetch = self._faiss.ntotal if doc_ids else min(n, self._faiss.ntotal)
        scores, idx = self._faiss.search(qvec, fetch)
        out: list[tuple[int, float]] = []
        for s, i in zip(scores[0], idx[0]):
            if i < 0 or (doc_ids and self._row_docs[i] not in doc_ids):
                continue
            out.append((self._row_ids[i], float(s)))
            if len(out) >= n:
                break
        return out

    def _keyword(self, query: str, n: int, doc_ids: set[str] | None) -> list[tuple[int, float]]:
        tokens = [t for t in re.findall(r"\w+", query.lower()) if t not in STOPWORDS]
        if not tokens:
            return []
        # Quote every token so user input can never be parsed as FTS5 syntax.
        match = " OR ".join(f'"{t}"' for t in dict.fromkeys(tokens))
        sql = (
            "SELECT c.id AS id, -bm25(chunks_fts) AS score FROM chunks_fts "
            "JOIN chunks c ON c.id = chunks_fts.rowid WHERE chunks_fts MATCH ?"
        )
        params: list = [match]
        if doc_ids:
            sql += f" AND c.doc_id IN ({','.join('?' * len(doc_ids))})"
            params += sorted(doc_ids)
        sql += " ORDER BY score DESC LIMIT ?"
        params.append(n)
        return [(r["id"], float(r["score"])) for r in self.db.execute(sql, params)]

    def _cosines(self, qvec: np.ndarray, ids: list[int]) -> dict[int, float]:
        if not ids:
            return {}
        marks = ",".join("?" * len(ids))
        rows = self.db.execute(f"SELECT id, embedding FROM chunks WHERE id IN ({marks})", ids)
        return {r["id"]: float(np.dot(np.frombuffer(r["embedding"], dtype=np.float32), qvec[0])) for r in rows}

    def _fetch_chunks(self, ids: list[int]) -> dict[int, Chunk]:
        if not ids:
            return {}
        marks = ",".join("?" * len(ids))
        rows = self.db.execute(f"SELECT {CHUNK_COLUMNS} FROM chunks WHERE id IN ({marks})", ids)
        return {r["id"]: Chunk(**{k: r[k] for k in r.keys() if k != "id"}) for r in rows}

    def search(
        self,
        query: str,
        k: int = 5,
        mode: Mode = "hybrid",
        doc_ids: list[str] | None = None,
        pool: int = 20,
    ) -> list[Hit]:
        """Hybrid = reciprocal-rank fusion of the semantic and keyword top-`pool` lists.
        RRF works on ranks, so it needs no score normalisation between cosine and BM25."""
        allowed = set(doc_ids) if doc_ids else None
        qvec = None
        if mode != "keyword":
            qvec = np.ascontiguousarray(self.embedder.embed_query(query), dtype=np.float32)
        sem = self._semantic(qvec, pool, allowed) if qvec is not None else []
        kw = self._keyword(query, pool, allowed) if mode != "semantic" else []
        sem_map, kw_map = dict(sem), dict(kw)

        if mode == "semantic":
            ranked = sem
        elif mode == "keyword":
            ranked = kw
        else:
            fused: dict[int, float] = defaultdict(float)
            for lst in (sem, kw):
                for rank, (cid, _) in enumerate(lst, start=1):
                    fused[cid] += 1.0 / (RRF_K + rank)
            ranked = sorted(fused.items(), key=lambda x: x[1], reverse=True)
        ranked = ranked[:k]

        if qvec is not None:  # cosine for every returned chunk (needed for abstention thresholds later)
            sem_map.update(self._cosines(qvec, [cid for cid, _ in ranked if cid not in sem_map]))
        chunks = self._fetch_chunks([cid for cid, _ in ranked])
        return [
            Hit(chunk=chunks[cid], score=score, semantic_score=sem_map.get(cid), keyword_score=kw_map.get(cid))
            for cid, score in ranked
        ]


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(prog="python -m docmind.index")
    ap.add_argument("--db", default=DEFAULT_DB)
    sub = ap.add_subparsers(dest="cmd", required=True)
    add = sub.add_parser("add", help="parse + index files")
    add.add_argument("paths", nargs="+")
    sub.add_parser("list", help="list indexed documents")
    dele = sub.add_parser("delete", help="remove a document by doc_id")
    dele.add_argument("doc_id")
    srch = sub.add_parser("search", help="retrieve chunks for a query")
    srch.add_argument("query")
    srch.add_argument("-k", type=int, default=5)
    srch.add_argument("--mode", choices=["hybrid", "semantic", "keyword"], default="hybrid")
    args = ap.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    from .embeddings import BgeEmbedder

    index = Index(BgeEmbedder(), args.db)

    if args.cmd == "add":
        for path in args.paths:
            try:
                doc_id, n = index.add_file(path)
            except ParseError as e:
                print(f"error: {e}")
                continue
            status = f"{n} chunks indexed" if n else "skipped (already indexed or no text)"
            print(f"{Path(path).name} [{doc_id}]: {status}")
    elif args.cmd == "list":
        for d in index.documents():
            print(f"{d['doc_id']}  {d['n_chunks']:>4} chunks  {d['doc_name']}")
    elif args.cmd == "delete":
        print("deleted" if index.delete_document(args.doc_id) else "no such document")
    else:
        for rank, h in enumerate(index.search(args.query, args.k, args.mode), start=1):
            sem = f"{h.semantic_score:.3f}" if h.semantic_score is not None else "-"
            kw = f"{h.keyword_score:.2f}" if h.keyword_score is not None else "-"
            print(f"\n#{rank} score={h.score:.4f} cos={sem} bm25={kw}  {h.chunk.citation}")
            print("   " + h.chunk.text[:220].replace("\n", " ") + ("..." if len(h.chunk.text) > 220 else ""))


if __name__ == "__main__":
    main()
'''
FILES['tests/test_chunking.py'] = r'''import pytest

from docmind.chunking import chunk_segments
from docmind.schemas import Segment


def seg(text, page=1):
    return Segment(doc_id="d1", doc_name="a.pdf", page=page, text=text)


def words(n, prefix="w"):
    return " ".join(f"{prefix}{i}" for i in range(n))


def test_long_paragraph_splits_with_overlap():
    chunks = chunk_segments([seg(words(500))], max_words=100, overlap_words=20)
    assert len(chunks) == 6
    assert all(c.word_count <= 100 for c in chunks)
    assert chunks[0].text.split()[-20:] == chunks[1].text.split()[:20]


def test_paragraphs_are_packed_up_to_limit():
    text = "\n\n".join(words(40, p) for p in "abc")
    chunks = chunk_segments([seg(text)], max_words=100)
    assert [c.word_count for c in chunks] == [80, 40]


def test_chunks_never_cross_pages():
    chunks = chunk_segments([seg(words(50, "a"), page=1), seg(words(50, "b"), page=2)])
    assert [c.page for c in chunks] == [1, 2]
    assert chunks[0].text.startswith("a0") and chunks[1].text.startswith("b0")


def test_tiny_and_empty_segments_are_dropped():
    chunks = chunk_segments([seg("12"), seg(""), seg(words(30), page=3)])
    assert len(chunks) == 1 and chunks[0].page == 3


def test_chunk_ids_unique_and_citation():
    chunks = chunk_segments([seg(words(500))], max_words=100, overlap_words=20)
    assert len({c.chunk_id for c in chunks}) == len(chunks)
    assert chunks[0].citation == "a.pdf, p. 1"

def test_overlap_must_be_smaller_than_max():
    with pytest.raises(ValueError):
        chunk_segments([seg(words(10))], max_words=10, overlap_words=10)
'''
FILES['tests/test_parsing.py'] = r'''import pymupdf
import pytest
from docx import Document

from docmind.parsing import ParseError, parse_document


def make_pdf(path, pages):
    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()


def test_pdf_pages_and_scanned_detection(tmp_path):
    p = tmp_path / "mixed.pdf"
    make_pdf(p, ["Invoice number 12345 issued by Acme Corp on 2024-01-05.", ""])
    segs = parse_document(p)
    assert [s.page for s in segs] == [1, 2]
    assert not segs[0].needs_ocr and "Acme Corp" in segs[0].text
    assert segs[1].needs_ocr and segs[1].text == ""


def test_docx_sections_and_tables(tmp_path):
    p = tmp_path / "contract.docx"
    d = Document()
    d.add_heading("Payment Terms", level=1)
    d.add_paragraph("Payment is due within 30 days of the invoice date.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Net"
    t.rows[0].cells[1].text = "30"
    d.add_heading("Termination", level=1)
    d.add_paragraph("Either party may terminate with 60 days written notice.")
    d.save(p)

    segs = parse_document(p)
    assert [s.section for s in segs] == ["Payment Terms", "Termination"]
    assert "Net | 30" in segs[0].text
    assert all(s.page is None for s in segs)


def test_unsupported_and_missing(tmp_path):
    f = tmp_path / "x.txt"
    f.write_text("hi")
    with pytest.raises(ParseError):
        parse_document(f)
    with pytest.raises(ParseError):
        parse_document(tmp_path / "nope.pdf")


def test_corrupt_pdf_raises_parse_error(tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"not a pdf")
    with pytest.raises(ParseError):
        parse_document(f)

def test_pdf_repeated_header_footer_removed(tmp_path):
    bodies = [
        "Revenue grew strongly across all regions during the year.",
        "The board approved the new dividend policy in March.",
        "Operating costs declined after the warehouse consolidation.",
        "Cash reserves remain sufficient for planned acquisitions.",
    ]
    p = tmp_path / "hf.pdf"
    pdf = pymupdf.open()
    for n, body in enumerate(bodies, start=1):
        page = pdf.new_page()
        page.insert_text((72, 40), "ACME Confidential Report 2024")
        page.insert_text((72, 200), body)
        page.insert_text((72, 800), f"Page {n} of 4")
    pdf.save(p)
    pdf.close()

    segs = parse_document(p)
    assert len(segs) == 4
    for seg, body in zip(segs, bodies):
        assert body in seg.text
        assert "ACME" not in seg.text and "Page" not in seg.text


def test_docx_without_headings_gets_paragraph_locator(tmp_path):
    p = tmp_path / "plain.docx"
    d = Document()
    d.add_paragraph("First paragraph of a document that has no headings at all.")
    d.add_paragraph("Second paragraph continues the same plain document text.")
    d.save(p)
    segs = parse_document(p)
    assert [s.section for s in segs] == ["paras 1-2"]
'''
FILES['tests/test_index.py'] = r'''import re

import numpy as np
import pytest

from docmind.index import Index
from docmind.schemas import Chunk


class FakeEmbedder:
    """Deterministic bag-of-words hashing embedder: no model download in tests."""

    DIM = 64

    def _vec(self, text: str) -> np.ndarray:
        v = np.zeros(self.DIM, dtype=np.float32)
        for tok in re.findall(r"\w+", text.lower()):
            v[hash(tok) % self.DIM] += 1.0
        n = np.linalg.norm(v)
        return v / n if n else v

    def embed_passages(self, texts):
        return np.stack([self._vec(t) for t in texts])

    def embed_query(self, query):
        return self._vec(query)[None, :]


TEXTS = [
    "Payment is due within 30 days of the invoice date.",
    "Either party may terminate this agreement with 60 days written notice.",
    "The invoice number is INV-2024-0042 issued by Acme Corp.",
    "The office is closed on public holidays.",
]


def make_chunks(doc_id="d1", name="a.pdf", texts=TEXTS):
    return [
        Chunk(chunk_id=f"{doc_id}:{i:04d}", doc_id=doc_id, doc_name=name, page=i + 1,
              chunk_index=i, text=t, word_count=len(t.split()))
        for i, t in enumerate(texts)
    ]


@pytest.fixture
def idx():
    index = Index(FakeEmbedder(), ":memory:")
    index.add_chunks(make_chunks())
    return index


def test_keyword_finds_exact_identifier(idx):
    hits = idx.search("INV-2024-0042", k=3, mode="keyword")
    assert hits[0].chunk.page == 3
    assert hits[0].keyword_score is not None


def test_semantic_finds_overlapping_text(idx):
    hits = idx.search("terminate agreement written notice", k=1, mode="semantic")
    assert hits[0].chunk.page == 2


def test_hybrid_fills_cosine_for_every_hit(idx):
    hits = idx.search("payment due days", k=3)
    assert hits[0].chunk.page == 1
    assert all(h.semantic_score is not None for h in hits)


def test_doc_filter(idx):
    idx.add_chunks(make_chunks("d2", "b.pdf", ["Payment terms are net 45 days for all invoices here."]))
    hits = idx.search("payment days", k=5, doc_ids=["d2"])
    assert hits and {h.chunk.doc_id for h in hits} == {"d2"}


def test_delete_removes_from_both_indexes(idx):
    assert idx.delete_document("d1")
    assert idx.search("payment", mode="keyword") == []
    assert idx.search("payment", mode="semantic") == []
    assert not idx.delete_document("d1")


def test_fts_syntax_in_query_does_not_crash(idx):
    idx.search('payment" OR ( NEAR', mode="keyword")


def test_empty_index_returns_nothing():
    assert Index(FakeEmbedder(), ":memory:").search("anything") == []
'''
FILES['requirements.txt'] = r'''pydantic>=2.5
pymupdf>=1.24.3
python-docx>=1.1
numpy
faiss-cpu
sentence-transformers
pytest>=8.0
'''
FILES[".gitignore"] = r'''.venv/
__pycache__/
.pytest_cache/
data/
*.gguf
'''

for rel, content in FILES.items():
    path = Path(rel)
    if rel == ".gitignore" and path.exists() and path.stat().st_size > 0:
        print(f"kept existing: {rel}")
        continue
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    print(f"wrote: {rel}")
print("\nDone. Next: pip install -r requirements.txt  ->  python -m pytest -v")