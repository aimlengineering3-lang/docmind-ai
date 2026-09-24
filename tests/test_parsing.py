import pymupdf
import pytest
from docx import Document

from docmind.parsing import ParseError, parse_document


def make_pdf(path, pages):
    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    pdf.save(path)
    pdf.close()


def test_pdf_pages_and_scanned_detection(tmp_path):
    p = tmp_path / "mixed.pdf"
    make_pdf(p, ["Invoice number 12345 issued by Acme Corp on 2024-01-05.", ""])
    segs = parse_document(p)
    assert [s.page for s in segs] == [1, 2]
    assert not segs[0].needs_ocr and "Acme Corp" in segs[0].text
    assert segs[1].needs_ocr and segs[1].text == ""


def test_docx_sections_and_tables(tmp_path):
    p = tmp_path / "contract.docx"
    d = Document()
    d.add_heading("Payment Terms", level=1)
    d.add_paragraph("Payment is due within 30 days of the invoice date.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Net"
    t.rows[0].cells[1].text = "30"
    d.add_heading("Termination", level=1)
    d.add_paragraph("Either party may terminate with 60 days written notice.")
    d.save(p)

    segs = parse_document(p)
    assert [s.section for s in segs] == ["Payment Terms", "Termination"]
    assert "Net | 30" in segs[0].text
    assert all(s.page is None for s in segs)


def test_unsupported_and_missing(tmp_path):
    f = tmp_path / "x.txt"
    f.write_text("hi")
    with pytest.raises(ParseError):
        parse_document(f)
    with pytest.raises(ParseError):
        parse_document(tmp_path / "nope.pdf")


def test_corrupt_pdf_raises_parse_error(tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"not a pdf")
    with pytest.raises(ParseError):
        parse_document(f)